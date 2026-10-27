"""Portable XLSX adapter; original workbooks are always read into a new export.

The standalone harness uses public openpyxl. Inside a Codex host that supplies an
artifact spreadsheet runtime, the agent may render/author through that runtime
using these same records. No private OpenAI package is required by this repo.
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

from .core import digest, file_hash, read_json, require
from .sync import validate_pairs

PAIR_HEADERS = ["配對編號", "同步依據", "A機卡別／原檔名", "B機卡別／原檔名", "內容／場景",
                "A片內入點", "A片內出點", "B片內入點", "B片內出點", "實拍重疊（秒）", "B－A偏移（秒）", "A播放倍率", "配對內容／同步錨點"]


def seconds(value):
    if isinstance(value, dt.timedelta):
        return value.total_seconds()
    if isinstance(value, dt.time):
        return value.hour * 3600 + value.minute * 60 + value.second + value.microsecond / 1e6
    if isinstance(value, (int, float)):
        return value * 86400
    parts = str(value).split(":")
    require(1 <= len(parts) <= 3, "invalid_timecode", "Expected Excel time or h:m:s timestamp")
    return sum(float(x) * 60 ** i for i, x in enumerate(reversed(parts)))


def tc(value):
    n = max(0, round(float(value) * 100))
    return f"{n // 360000:02}:{n // 6000 % 60:02}:{n // 100 % 60:02}.{n % 100:02}"


def _map(ctx, stage):
    data = ctx.stage(stage, required=False)
    return {x["clip_id"]: x for x in data.get("clips", [])} if data else {}


def unique_sheet(wb, name):
    candidate, n = name, 2
    while candidate in wb.sheetnames:
        candidate = name[:27] + str(n); n += 1
    return wb.create_sheet(candidate)


def literal(cell, value):
    cell.value = value
    # Model text and media filenames are data, never spreadsheet formulas.
    if isinstance(value, str):
        cell.data_type = "s"


def format_table(sheet, header_row, last_row, widths, table_name):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter
    for i, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    for row in sheet.iter_rows(min_row=header_row, max_row=last_row, max_col=len(widths)):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in sheet[header_row][:len(widths)]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="196782")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[header_row].height = 34
    sheet.freeze_panes = "C" + str(header_row + 1)
    sheet.sheet_view.showGridLines = False
    if last_row > header_row:
        table = Table(displayName=table_name, ref=f"A{header_row}:{get_column_letter(len(widths))}{last_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)


def export(ctx, output, template=None):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    output = Path(output).expanduser().resolve()
    require(not output.exists(), "output_exists", "Refusing to overwrite an existing workbook")
    require(output.suffix.lower() == ".xlsx", "unsupported_workbook", "Use an .xlsx output")
    clips = ctx.stage("inventory")["clips"]
    require(len({(c["card"], c["file"]) for c in clips}) == len(clips), "ambiguous_filename", "Workbook card rows require unique basenames within each card", "Split duplicate-name subfolders into distinct card IDs.")
    reviews, transcripts, environment, local = (_map(ctx, n) for n in ("reviews", "transcripts", "environment", "local_vision"))
    pair_stage = ctx.stage("pairs", required=False); pairs = pair_stage["pairs"] if pair_stage else []
    if pairs:
        validate_pairs(pairs, clips)
    by_id = {c["id"]: c for c in clips}; counterparts = {}
    for row in pairs:
        for own, other in ((row["a_id"], row["b_id"]), (row["b_id"], row["a_id"])):
            counterparts.setdefault(own, []).append(row["group"] + " " + other)
    original = {}
    if template:
        template = Path(template).expanduser().resolve()
        require(template.suffix.lower() == ".xlsx", "unsupported_workbook", "Only .xlsx templates are supported; macros are not rewritten")
        wb = openpyxl.load_workbook(template)
        for s in wb:
            original[s.title] = {c.coordinate: c.value for row in s for c in row if c.value is not None}
    else:
        wb = openpyxl.Workbook(); wb.active.title = "總覽"
        summary = wb.active
        summary["A1"] = "飲食紀錄片場記"
        summary["A2"] = "原檔時間碼；場記內容與原音轉錄分開保存。"
        summary["A3"] = f"{len(clips)} 支素材；{len(reviews)} 支已有畫面核對紀錄；{len(pairs)} 組 AB 配對。"
        summary.column_dimensions["A"].width = 115
    token = digest(str(output))[:8]
    for card in ctx.cards:
        card_clips = [c for c in clips if c["card"] == card]
        if not card_clips:
            continue
        if card in wb.sheetnames and template:
            s = wb[card]; start_col = s.max_column + 1
            mapping = {}
            for row in s:
                matches = [c for c in row if isinstance(c.value, str) and c.value.lower().endswith((".mp4", ".mov", ".mxf", ".mts"))]
                for cell in matches:
                    mapping.setdefault(Path(cell.value).name, []).append(cell.row)
            require(all(len(mapping.get(c["file"], [])) == 1 for c in card_clips), "template_mapping_ambiguous",
                    "Could not map every source file to exactly one existing card row", "Use a template with one exact filename per card row, or create a new workbook.", card=card)
            header = 7
        else:
            s = unique_sheet(wb, card); start_col = 5; header = 7
            literal(s["A1"], card + " 場記")
            for i, text in enumerate(["序號", "原檔名", "原檔時長", "播放幀率"], 1):
                literal(s.cell(header, i), text)
            mapping = {}
            for index, c in enumerate(card_clips, 1):
                row = header + index; mapping[c["file"]] = [row]
                for col, value in enumerate([index, c["file"], c["duration"] / 86400, c["fps"]], 1):
                    literal(s.cell(row, col), value)
                s.cell(row, 3).number_format = "[m]:ss.000"
                s.cell(row, 4).number_format = "0.000"
        headers = ["景別", "畫面與工序", "環境音／聲響", "對白／片內時間碼", "AB對機素材", "辨識依據"]
        for i, text in enumerate(headers, start_col):
            literal(s.cell(header, i), text)
            s.column_dimensions[get_column_letter(i)].width = [16, 70, 60, 90, 62, 34][i - start_col]
            s.cell(header, i).font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            s.cell(header, i).fill = PatternFill("solid", fgColor="196782")
        for c in card_clips:
            row = mapping[c["file"]][0]; review = reviews.get(c["id"], {})
            draft = local.get(c["id"], {}); tr = transcripts.get(c["id"], {})
            chosen = review["meaningful_dialogue"] if "meaningful_dialogue" in review else tr.get("segments", [])
            dialogue = "\n".join(f"{tc(x['start'])}–{tc(x['end'])} {x['text']}" for x in chosen[:5])
            if len(chosen) > 5:
                dialogue += f"\n另有 {len(chosen) - 5} 段，詳逐段對白。"
            env = environment.get(c["id"], {}); sounds = []
            for window in env.get("windows", []):
                if window.get("description"):
                    sounds.append(window["description"])
                else:
                    sounds.extend(x["label"] for x in window.get("tags", [])[:3])
            sounds = list(dict.fromkeys(sounds))
            basis = "Codex畫面核對" if review.get("review_mode") == "codex_native_vision" else "人工畫面核對" if review else "本機模型" if draft else "素材中繼資料"
            if c["ratio"] != 1:
                basis += f"；慢動作 {c['ratio']:g} 倍"
            values = [review.get("shot_size", ""), review.get("scene_description", draft.get("description", "")),
                      ("模型候選：\n" if sounds else "") + "\n".join(sounds[:8]) + ("\n詳逐段環境音。" if len(sounds) > 8 else ""), dialogue,
                      "\n".join(counterparts.get(c["id"], [])), basis]
            for col, value in enumerate(values, start_col):
                cell = s.cell(row, col); literal(cell, value)
                cell.font = Font(name="Arial", size=10); cell.alignment = Alignment(vertical="top", wrap_text=True)
            s.row_dimensions[row].height = max(s.row_dimensions[row].height or 15, min(409, 18 * max(4, *(str(x).count("\n") + 2 for x in values))))
        s.freeze_panes = s.freeze_panes or "C8"
        s.sheet_view.showGridLines = False
        if not template or card not in original:
            format_table(s, header, header + len(card_clips), [8, 34, 16, 12, 16, 70, 60, 90, 62, 34], "Card" + digest(card + token)[:12])
    if pairs:
        sheet = unique_sheet(wb, "AB機同步對照")
        literal(sheet["A1"], "AB機同步對照")
        literal(sheet["A2"], "A 下 B 上。偏移定義：B 片內時間 = A 片內時間 + 偏移（僅原音波形列）。")
        for col, value in enumerate(PAIR_HEADERS, 1):
            literal(sheet.cell(7, col), value)
        for index, pair in enumerate(pairs, 8):
            review = reviews.get(pair["a_id"], {})
            description = pair.get("subject") or review.get("scene_description", "")
            basis = "多段原音波形吻合" if pair["basis"] == "原音波形" else "校正相機時鐘與實拍時長後配對"
            row = [pair["group"], pair["basis"], pair["a_id"], pair["b_id"], description,
                   pair["a_in"] / 86400, pair["a_out"] / 86400, pair["b_in"] / 86400, pair["b_out"] / 86400,
                   pair["duration"], pair.get("lag"), pair["ratio"], basis]
            for col, value in enumerate(row, 1):
                literal(sheet.cell(index, col), value)
            for col in range(6, 10):
                sheet.cell(index, col).number_format = "[m]:ss.000"
            sheet.cell(index, 10).number_format = "0.000"
            sheet.cell(index, 11).number_format = "+0.000;-0.000;0.000"
            sheet.cell(index, 12).number_format = '0.0"x"'
            sheet.row_dimensions[index].height = 65
        format_table(sheet, 7, 7 + len(pairs), [10, 24, 47, 47, 66, 15, 15, 15, 15, 16, 17, 14, 52], "ABPairs" + token)
        pair_sheet = sheet.title
    else:
        pair_sheet = None
    detail = unique_sheet(wb, "逐段對白")
    for col, value in enumerate(["素材ID", "片內起點", "片內終點", "聲道", "Whisper原稿", "有意義對白／修訂", "依據"], 1):
        literal(detail.cell(1, col), value)
    row = 2
    for c in clips:
        for segment in transcripts.get(c["id"], {}).get("segments", []):
            values = [c["id"], segment["start"] / 86400, segment["end"] / 86400, segment.get("channel"), segment["text"], "", "Whisper"]
            for col, value in enumerate(values, 1):
                literal(detail.cell(row, col), value)
            detail.cell(row, 2).number_format = detail.cell(row, 3).number_format = "[m]:ss.000"; row += 1
        for segment in reviews.get(c["id"], {}).get("meaningful_dialogue", []):
            values = [c["id"], segment["start"] / 86400, segment["end"] / 86400, segment.get("channel"), "", segment["text"], segment["basis"]]
            for col, value in enumerate(values, 1):
                literal(detail.cell(row, col), value)
            detail.cell(row, 2).number_format = detail.cell(row, 3).number_format = "[m]:ss.000"; row += 1
    format_table(detail, 1, max(1, row - 1), [52, 16, 16, 10, 92, 92, 24], "Dialogue" + token)
    env_sheet = unique_sheet(wb, "逐段環境音")
    for col, value in enumerate(["素材ID", "片內起點", "片內終點", "模型觀察／候選聲響", "模型"], 1):
        literal(env_sheet.cell(1, col), value)
    row = 2
    for c in clips:
        env = environment.get(c["id"], {})
        for window in env.get("windows", []):
            description = window.get("description") or "；".join(f"{t['label']} ({t['score']:.3f})" for t in window.get("tags", []))
            for col, value in enumerate([c["id"], window["start"] / 86400, window["end"] / 86400, description, env.get("model")], 1):
                literal(env_sheet.cell(row, col), value)
            env_sheet.cell(row, 2).number_format = env_sheet.cell(row, 3).number_format = "[m]:ss.000"; row += 1
    format_table(env_sheet, 1, max(1, row - 1), [52, 16, 16, 115, 60], "Environment" + token)
    for name, values in original.items():
        for coord, value in values.items():
            require(wb[name][coord].value == value, "template_changed", "An existing workbook cell changed", sheet=name, cell=coord)
    output.parent.mkdir(parents=True, exist_ok=True); wb.save(output)
    reopened = openpyxl.load_workbook(output)
    for name, values in original.items():
        for coord, value in values.items():
            require(reopened[name][coord].value == value, "export_changed", "Original content changed on XLSX round trip", sheet=name, cell=coord)
    return {"output": str(output), "bytes": output.stat().st_size, "sha256": file_hash(output), "pair_sheet": pair_sheet,
            "source_cells_preserved": sum(len(v) for v in original.values()), "visual_review_required": True,
            "note": "Open or render the workbook to check layout before delivery. Portable export does not claim pixel-level verification."}


def import_pairs(ctx, path, sheet="AB機同步對照"):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    inventory = ctx.stage("inventory")["clips"]
    known = {c["id"] for c in inventory}
    def canonical_id(value):
        if value in known:
            return value
        require(isinstance(value, str) and "/" in value, "unknown_clip", "Expected CARD/filename or a full inventory ID", value=value)
        card, remainder = value.split("/", 1)
        matches = [c["id"] for c in inventory if c["card"] == card and c["file"] == Path(remainder).name]
        require(len(matches) == 1, "ambiguous_clip", "Cannot map legacy CARD/filename to exactly one clip", value=value, matches=matches)
        return matches[0]
    require(sheet in wb.sheetnames, "sheet_missing", "AB pairing sheet was not found", available=wb.sheetnames)
    ws = wb[sheet]; rows = []
    for values in ws.iter_rows(min_row=8, values_only=True):
        if not values or not values[0]:
            continue
        require(len(values) >= 13, "schema_mismatch", "Pair sheet must use the documented 13-column layout")
        group, basis, aid, bid, subject, ain, aout, bin_, bout, duration, lag, ratio, note = values[:13]
        rows.append({"group": group, "basis": basis, "a_id": canonical_id(aid), "b_id": canonical_id(bid), "subject": subject or "",
                     "a_in": seconds(ain), "a_out": seconds(aout), "b_in": seconds(bin_), "b_out": seconds(bout),
                     "duration": float(duration), "lag": float(lag) if lag is not None else None, "ratio": float(ratio), "note": note or ""})
    require(rows, "no_pairs", "Pair sheet contains no rows")
    checked = validate_pairs(rows, ctx.stage("inventory")["clips"])
    return {**ctx.publish("pairs", {"pairs": rows, "source_workbook": str(Path(path).resolve()), "source_sha256": file_hash(path)}), **checked}
