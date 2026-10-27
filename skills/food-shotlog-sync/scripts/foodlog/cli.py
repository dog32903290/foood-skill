"""Stable command grammar and one JSON document per machine-readable call."""
from __future__ import annotations

import argparse
import json
from pathlib import Path
import re
import sys

from . import __version__
from .core import Context, HarnessError, PROTOCOL, file_hash, read_json, require, source_unchanged, status, write_json


class Parser(argparse.ArgumentParser):
    def error(self, message):
        raise HarnessError("invalid_arguments", message, "Run foodlog --help or the command's --help.")


def parser():
    p = Parser(prog="foodlog", description="Local shot logs, native Codex vision evidence, and Resolve A/B timelines")
    p.add_argument("--json", action="store_true", help="Emit one machine-readable response")
    p.add_argument("--config", help="Local project configuration; never committed with footage")
    p.add_argument("--version", action="version", version=__version__)
    commands = p.add_subparsers(dest="command", required=True)
    commands.add_parser("status", help="Read-only backend and session probe")
    init = commands.add_parser("init", help="Create a new local project configuration")
    init.add_argument("--source-root", required=True); init.add_argument("--work-dir", required=True)
    init.add_argument("--output", default="foodlog.local.json")
    init.add_argument("--card", action="append", help="CARD:CAMERA:relative/path; repeat for each card")
    inv = commands.add_parser("inventory", help="Read video metadata and Sony XML without changing sources")
    inv.add_argument("--limit", type=int, help="Optional explicit whole-inventory limit for a small job")
    ls = commands.add_parser("list", help="List stable clip IDs or available command capabilities")
    ls.add_argument("kind", choices=["clips", "capabilities"]); ls.add_argument("--limit", type=int, default=30)
    audio = commands.add_parser("audio", help="Prepare measured multi-channel 16 kHz PCM caches")
    audio.add_argument("action", choices=["prepare"])
    fetch = commands.add_parser("model-download", help="Explicitly fetch one allowlisted open model")
    fetch.add_argument("backend", choices=["faster-whisper", "mlx-whisper", "ast"]); fetch.add_argument("name")
    asr = commands.add_parser("asr", help="Offline Whisper transcription; exact chunk offsets, faster-whisper VAD")
    asr.add_argument("--clip-id", action="append")
    env = commands.add_parser("environment", help="Local environmental audio classification or audio-model description")
    env.add_argument("--backend", choices=["ast", "local-audio"], default="ast"); env.add_argument("--clip-id", action="append")
    commands.add_parser("polish", help="Local text-model suggestions; never overwrite the Whisper original")
    review = commands.add_parser("review", help="Real frame bundles, local vision drafts, and native Codex review records")
    rs = review.add_subparsers(dest="action", required=True)
    pack = rs.add_parser("pack"); pack.add_argument("--clip-id", action="append"); pack.add_argument("--frames", type=int, default=4)
    pack.add_argument("--batch-size", type=int, default=12)
    rs.add_parser("local")
    imp = rs.add_parser("import"); imp.add_argument("path")
    sync = commands.add_parser("sync", help="Correlate microphones and calibrate camera clocks for silent clips")
    sync.add_argument("--clock-bias-seconds", type=float, help="A clock minus B clock; omit to calibrate from waveforms")
    sheet = commands.add_parser("sheet", help="Portable XLSX export and authoritative pairing-table import")
    ss = sheet.add_subparsers(dest="action", required=True)
    ex = ss.add_parser("export"); ex.add_argument("--output", required=True); ex.add_argument("--template")
    si = ss.add_parser("import-pairs"); si.add_argument("path"); si.add_argument("--sheet", default="AB機同步對照")
    inspect = ss.add_parser("inspect"); inspect.add_argument("path")
    resolve = commands.add_parser("resolve", help="Read, plan, apply, inspect, capture and export using the real Resolve SDK")
    ds = resolve.add_subparsers(dest="action", required=True)
    ds.add_parser("status"); ds.add_parser("plan"); ds.add_parser("inspect")
    apply = ds.add_parser("apply"); apply.add_argument("--execute", action="store_true"); apply.add_argument("--resume", action="store_true")
    export = ds.add_parser("export"); export.add_argument("--directory", required=True)
    capture = ds.add_parser("capture"); capture.add_argument("--output", required=True); capture.add_argument("--window-id", type=int)
    ui = ds.add_parser("ui-review"); ui.add_argument("path")
    commands.add_parser("validate", help="Validate source fingerprints, timestamp bounds and pairing invariants")
    commands.add_parser("logs", help="Locate session head, immutable bundles, and command trajectory")
    return p


def initialize(args):
    source = Path(args.source_root).expanduser().resolve(); work = Path(args.work_dir).expanduser().resolve()
    require(source.is_dir(), "source_missing", "Source root does not exist")
    cards = {}
    if args.card:
        for spec in args.card:
            parts = spec.split(":", 2)
            require(len(parts) == 3, "invalid_card", "Card syntax is CARD:CAMERA:relative/path")
            name, camera, path = parts
            require(camera in ("A", "B"), "invalid_camera", "Camera must be A or B")
            cards[name] = {"camera": camera, "path": path}
    else:
        cards = {p.name: {"camera": p.name[0].upper(), "path": p.name} for p in sorted(source.iterdir())
                 if p.is_dir() and re.fullmatch(r"[ABab]\d+", p.name)}
    require(cards, "cards_missing", "No A/B cards were found", "Pass --card CARD:CAMERA:relative/path or edit the example configuration.")
    value = {"schema_version": 1, "source_root": str(source), "work_dir": str(work), "cards": cards,
             "asr": {"backend": "faster-whisper", "model": "small", "language": "zh", "device": "cpu", "compute_type": "int8", "chunk_seconds": 60},
             "sync": {"max_waveform_comparisons": 10000, "clock_seed_pairs": 3, "clock_uncertainty_seconds": 2},
             "resolve": {"fps": "30000/1001", "fps_setting": "29.97", "width": 1920, "height": 1080, "gap_frames": 30,
                         "timeline_name": "AB雙機校對_A下B上", "start_timecode": "01:00:00:00"}}
    output = Path(args.output).expanduser().resolve()
    require(not output.exists(), "output_exists", "Refusing to overwrite an existing configuration")
    # Validate boundaries before committing the new config to its final path.
    for name, card in cards.items():
        root = (source / card["path"]).resolve()
        require(root.is_dir() and root.is_relative_to(source), "invalid_card", "Card must be a directory inside source_root", card=name)
        require(not work.is_relative_to(root), "unsafe_output", "Work directory cannot be inside a source card", card=name)
    write_json(output, value)
    return {"config": str(output), "cards": list(cards), "work_dir": str(work)}


def validate(ctx):
    from .sync import validate_pairs
    clips = ctx.stage("inventory")["clips"]
    for row in clips:
        source_unchanged(row)
    report = {"source_files_unchanged": len(clips), "source_ids_unique": len({c["id"] for c in clips}) == len(clips)}
    require(report["source_ids_unique"], "duplicate_id", "Inventory contains duplicate source IDs")
    pairs = ctx.stage("pairs", required=False)
    if pairs:
        report["pairs"] = validate_pairs(pairs["pairs"], clips)
    by_id = {c["id"]: c for c in clips}
    transcripts = ctx.stage("transcripts", required=False)
    for clip in transcripts.get("clips", []) if transcripts else []:
        require(clip["clip_id"] in by_id and clip["fingerprint"] == by_id[clip["clip_id"]]["fingerprint"], "stale_transcript", "Transcript is not tied to the current inventory")
        for s in clip["segments"]:
            require(0 <= s["start"] < s["end"] <= by_id[clip["clip_id"]]["duration"] + .1,
                    "invalid_timecode", "Transcript lies outside the source duration", clip_id=clip["clip_id"])
    reviews = ctx.stage("reviews", required=False)
    for review in reviews.get("clips", []) if reviews else []:
        require(review["clip_id"] in by_id and review["fingerprint"] == by_id[review["clip_id"]]["fingerprint"], "stale_review", "Visual review is not tied to the current inventory")
        for asset in review["evidence"]:
            require(file_hash(asset["path"]) == asset["sha256"], "evidence_changed", "A reviewed image changed")
    report["visually_reviewed_clips"] = len(reviews["clips"]) if reviews else 0
    report["semantic_verification"] = "Requires the documented native Codex/human image and audio review; numerical validation does not verify the words or food identity."
    return report


def dispatch(args):
    if args.command == "init":
        return initialize(args)
    if args.command == "model-download":
        from .recognition import download_model
        return download_model(args.backend, args.name)
    if args.command == "list" and args.kind == "capabilities":
        return {"commands": ["inventory", "audio prepare", "asr", "environment", "polish", "review pack", "review local", "review import", "sync", "sheet export", "sheet import-pairs", "resolve plan", "resolve apply --execute", "resolve inspect", "resolve capture", "resolve ui-review", "resolve export", "validate", "logs"],
                "native_codex_vision": "Performed by the host agent after actually opening frame images; no private model API is simulated.",
                "default_layout": "A on V1 below B on V2", "source_mutation": False}
    if args.command == "sheet" and args.action == "inspect":
        import openpyxl
        wb = openpyxl.load_workbook(args.path, read_only=True, data_only=True)
        return {"path": str(Path(args.path).resolve()), "sha256": file_hash(args.path),
                "sheets": [{"name": s.title, "rows": s.max_row, "columns": s.max_column,
                            "first_rows": [[str(v) if v is not None else None for v in row] for row in s.iter_rows(min_row=1, max_row=7, max_col=13, values_only=True)]} for s in wb]}
    if args.command == "status" and not args.config:
        return status()
    if args.command == "resolve" and args.action == "status":
        from .resolve import status as resolve_status
        return resolve_status(Context(args.config) if args.config else None)
    require(args.config, "config_required", "Pass --config before the command")
    ctx = Context(args.config)
    if args.command == "status":
        return status(ctx)
    if args.command == "logs":
        return {"session": str(ctx.session_path), "trajectory": str(ctx.work / "trajectory.jsonl"), "head": ctx.head()}
    if args.command == "list":
        rows = ctx.stage("inventory")["clips"]
        return {"total": len(rows), "clips": [{k: c[k] for k in ("id", "camera", "duration", "ratio")} for c in rows[:args.limit]]}
    if args.command == "validate":
        return validate(ctx)
    if args.command == "resolve" and args.action == "inspect":
        from .resolve import inspect
        return inspect(ctx)
    with ctx.lock():
        if args.command == "inventory":
            from .media import inventory
            return inventory(ctx, args.limit)
        if args.command == "audio":
            from .media import audio_stage
            return audio_stage(ctx)
        if args.command == "asr":
            from .recognition import transcribe
            return transcribe(ctx, args.clip_id)
        if args.command == "environment":
            from .recognition import environment
            return environment(ctx, args.backend, args.clip_id)
        if args.command == "polish":
            from .recognition import polish
            return polish(ctx)
        if args.command == "review":
            from . import review
            if args.action == "pack":
                return review.make_pack(ctx, args.clip_id, args.frames, args.batch_size)
            if args.action == "local":
                return review.local_vision(ctx)
            return review.import_review(ctx, args.path)
        if args.command == "sync":
            from .sync import pair
            return pair(ctx, args.clock_bias_seconds)
        if args.command == "sheet":
            from . import workbook
            if args.action == "export":
                return workbook.export(ctx, args.output, args.template)
            return workbook.import_pairs(ctx, args.path, args.sheet)
        if args.command == "resolve":
            from . import resolve
            if args.action == "plan":
                return resolve.make_plan(ctx)
            if args.action == "apply":
                return resolve.apply(ctx, execute=args.execute, resume=args.resume)
            if args.action == "export":
                return resolve.export(ctx, args.directory)
            if args.action == "capture":
                return resolve.capture(ctx, args.output, args.window_id)
            if args.action == "ui-review":
                return resolve.ui_review(ctx, args.path)
    raise HarnessError("unknown_command", "Command has no implementation")


def main():
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8")
    try:
        args = parser().parse_args()
        result = dispatch(args)
        print(json.dumps({"ok": True, "protocol": PROTOCOL, "version": __version__, "command": args.command,
                          "result": result}, ensure_ascii=False, indent=None if args.json else 2, allow_nan=False))
    except HarnessError as exc:
        print(json.dumps({"ok": False, "protocol": PROTOCOL, "error": {"type": exc.kind, "message": str(exc),
                          "repair_hint": exc.hint}, "context": exc.context}, ensure_ascii=False))
        raise SystemExit(2)
    except Exception as exc:
        print(json.dumps({"ok": False, "protocol": PROTOCOL, "error": {"type": "unexpected_backend_error", "message": str(exc),
                          "repair_hint": "Inspect the named backend/input and session trajectory. No fake successful artifact was generated."}}, ensure_ascii=False))
        raise SystemExit(1)


if __name__ == "__main__":
    main()
